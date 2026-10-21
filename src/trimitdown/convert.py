"""Чистая конверсия документа в markdown.

Ни FastAPI, ни HTTP, ни архива: на вход путь или байты, на выходе текст и
счётчики. Веб-обёртки живут в core/converter.py и зовут этот модуль. Граница
закрыта тестом (tests/test_convert_pure.py::test_pure_module_does_not_import_fastapi):
импорт fastapi отсюда утяжелил бы CLI-пакет на весь веб-стек, ничего при этом не
сломав внешне -- то есть незаметно.
"""
import codecs
import os
import re
import tempfile
from dataclasses import dataclass
from importlib import resources
from pathlib import Path

import tiktoken

from pptx import Presentation

from trimitdown.pdf_route import PageRoute, render_markdown, route_pdf

# Кэш едет внутри пакета как package data. Path(__file__) сработал бы для обычной
# установки, но не переживает упаковку; resources.files находит его и там, и в
# рабочем дереве. Без этого tiktoken при первом счёте полезет в сеть -- в офлайне
# это отказ, в оплаченном туннеле сюрприз.
os.environ.setdefault(
    "TIKTOKEN_CACHE_DIR", str(resources.files("trimitdown") / "tiktoken_cache")
)

TOKENS_PER_UNIT_ESTIMATE = 2250  # midpoint of the documented 1500-3000 tokens/page vision-estimate
                                  # range (Anthropic docs) — used for the PDF/PPTX before/after
                                  # comparison only. The frontend never renders a negative saving:
                                  # if a dense page's extracted text still exceeds this estimate,
                                  # it shows the raw result-token count with no percentage.

_encoding = None


class ConversionError(Exception):
    """Документ не удалось сконвертировать.

    Своё исключение, а не HTTPException: ядро не знает про HTTP. Веб-слой ловит
    его и переводит в 422, CLI -- в ненулевой код возврата.
    """


@dataclass(frozen=True)
class ConversionResult:
    text: str
    tokens_after: int | None
    tokens_before: int | None
    unit: str | None
    units: int | None
    # PDFs only: how each page was routed (accept / substitute / scan / corrupt
    # / blank) and why. Last and defaulted, so the web layer, which builds
    # nothing from it, keeps working unchanged.
    pages: tuple[PageRoute, ...] | None = None


def _get_encoding():
    global _encoding
    if _encoding is None:
        _encoding = tiktoken.get_encoding("cl100k_base")
    return _encoding


def count_tokens(text: str) -> int:
    return len(_get_encoding().encode(text))


def safe_stem(name: str) -> str:
    # Untrusted filenames may carry either separator regardless of host OS --
    # this same function sanitizes uploads on the Linux-hosted Docker server
    # and on the Windows/macOS desktop app. Path().stem's separator and root
    # parsing differ by platform (Windows treats "\" as a separator and
    # collapses a leading "///" into a UNC-style prefix; POSIX does neither),
    # so a name built from Path() here would sanitize differently depending
    # on which OS the server happens to run on. Reduce explicitly instead.
    last = re.split(r"[/\\]", name)[-1]
    stem = last.rsplit(".", 1)[0]
    stem = re.sub(r"[^\w\-. ]", "_", stem, flags=re.UNICODE).strip()
    return stem or "file"


def _count_pdf_pages(path: Path) -> int:
    import pypdfium2

    document = pypdfium2.PdfDocument(str(path))
    try:
        return len(document)
    finally:
        document.close()


def _count_pptx_slides(path: Path) -> int:
    return len(Presentation(path).slides)


def _estimate_before_tokens(suffix: str, path: str) -> tuple[int | None, str | None, int | None]:
    try:
        if suffix == ".pdf":
            units = _count_pdf_pages(Path(path))
            return units * TOKENS_PER_UNIT_ESTIMATE, "page", units
        if suffix == ".pptx":
            units = _count_pptx_slides(Path(path))
            return units * TOKENS_PER_UNIT_ESTIMATE, "slide", units
    except Exception:
        pass
    return None, None, None


_NEEDS_OFFICE_EXTRA = (
    "this format needs the office extra: pip install \"trimitdown[office]\" "
    "(PDF needs no extra)"
)

# Everything the office extra reads, named in the refusal for an unknown file so
# the message says what to try instead of only what failed.
_SUPPORTED_NON_PDF = (
    "DOCX, DOC, RTF, ODT, EPUB, XLSX, XLS, ODS, CSV, PPTX, PPT, ODP, HTML, TXT, MD"
)

_TEXT_SUFFIXES = frozenset({".txt", ".md"})
_HTML_SUFFIXES = frozenset({".html", ".htm"})


def _convert_other(path: Path, suffix: str, head: bytes) -> str:
    """Everything that is not a PDF: plain text, HTML, or the office formats.

    Plain text and HTML are handled here because anydoc reads neither, and
    markitdown, which did, passed `.txt` and `.md` through unchanged; dropping
    them would have been a silent regression for anyone piping notes through
    the CLI. Everything else is anydoc's.

    HTML is recognised by suffix or by content. By content because a page saved
    as "page.pdf" has to keep converting -- the file is no PDF, so it must not
    be refused for the name it carries. Text suffixes are never sniffed: a
    markdown file that quotes an <html> tag is still markdown.
    """
    if suffix in _TEXT_SUFFIXES:
        return _decode_text(path.read_bytes(), path.name)
    if suffix in _HTML_SUFFIXES or _sniffs_as_html(head):
        return _html_to_markdown(_decode_html(path.read_bytes()))
    return _anydoc_to_markdown(path)


def _sniffs_as_html(head: bytes) -> bool:
    if head.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        text = head.decode("utf-16", errors="ignore")
    elif b"\x00" in head:
        # A zip or OLE header (DOCX, XLSX, DOC...) is full of NUL bytes; page
        # markup is not. Without this a binary container could match by chance.
        return False
    else:
        text = head.decode("latin-1")
    text = text.lower()
    return "<!doctype html" in text or "<html" in text


def _bom_encoding(data: bytes) -> str | None:
    if data.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"  # the codec reads the BOM and picks the byte order itself
    return None


def _decode_text(data: bytes, name: str) -> str:
    """Decode a plain-text file: a BOM if there is one, otherwise strict UTF-8.

    There is no guessing. A text file in a legacy Cyrillic code page is also
    valid bytes in other legacy code pages, so any guess would turn it into
    confident garbage that reaches an LLM unflagged -- the same reason the OCR
    path never falls back to English. A file that is not UTF-8 is refused,
    loudly, with the remedy in the message.
    """
    encoding = _bom_encoding(data) or "utf-8"
    try:
        return data.decode(encoding)
    except UnicodeDecodeError:
        raise ConversionError(
            f"{name} is not valid UTF-8 text; re-save it as UTF-8 (no encoding is guessed)"
        ) from None


_META_CHARSET = re.compile(rb"""<meta[^>]+charset\s*=\s*["']?\s*([A-Za-z0-9_.:\-]+)""", re.IGNORECASE)


def _decode_html(data: bytes) -> str:
    """Decode an HTML file the way HTML5 does, in this order.

    1. A BOM wins.
    2. A `<meta charset>` (or the http-equiv form) in the first 4 KiB, if the
       label is one Python knows. An unknown label is ignored, not fatal.
    3. UTF-8, if the bytes are valid UTF-8.
    4. windows-1252, the HTML5 default for an undeclared legacy page. Unlike
       plain text, HTML has a specified fallback, and old pages in it are
       common enough that refusing them would be unhelpful. Bytes the code page
       leaves undefined become U+FFFD instead of failing the conversion.
    """
    bom = _bom_encoding(data)
    if bom:
        return data.decode(bom, errors="replace")
    declared = _META_CHARSET.search(data[:4096])
    if declared:
        try:
            label = codecs.lookup(declared.group(1).decode("ascii")).name
        except LookupError:
            label = None
        if label:
            return data.decode(label, errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _html_to_markdown(html: str) -> str:
    # Lazy for the same reason as anydoc: a PDF must not pay for either import,
    # and the base install has neither.
    try:
        import markdownify
        from bs4 import BeautifulSoup
    except ImportError:
        raise ConversionError(_NEEDS_OFFICE_EXTRA) from None

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style"]):
        tag.extract()
    # Links that are not web or file links (javascript:, mailto:, tel:) keep
    # their words and lose the target: the target is noise in pasted text.
    for link in soup.find_all("a", href=True):
        scheme = link["href"].partition(":")[0].strip().lower()
        if ":" in link["href"] and scheme not in ("http", "https", "file"):
            del link["href"]
    # An inline data: URI is the picture's bytes in base64 -- thousands of
    # tokens that say nothing. Keep the alt text and the media type.
    for image in soup.find_all("img", src=True):
        if image["src"][:5].lower() == "data:":
            image["src"] = image["src"].split(",")[0] + "..."

    body = soup.find("body")
    root = soup if body is None else body
    try:
        text = markdownify.MarkdownConverter(heading_style=markdownify.ATX).convert_soup(root)
    except RecursionError:
        # markdownify walks the tree recursively; a pathologically nested page
        # exhausts the stack. The words are worth more than the markup.
        text = root.get_text("\n", strip=True)
    return text.strip()


def _describe_anydoc_error(error: Exception) -> str:
    """One human line per failure kind; anydoc's own text is the detail."""
    import anydoc

    detail = str(error)
    if isinstance(error, anydoc.UnsupportedError):
        return (
            "unsupported file type: the format was not recognised from its content "
            f"or extension (supported: {_SUPPORTED_NON_PDF}, PDF)"
        )
    if isinstance(error, anydoc.EncryptedError):
        return "the document is encrypted or password-protected; remove the protection and convert again"
    if isinstance(error, anydoc.NeedsOcrError):
        return (
            "the document is scanned or image-only and needs OCR, which this "
            "converter does not do (a PDF with a .pdf name goes through the PDF "
            "route, which hands scans to review)"
        )
    if isinstance(error, anydoc.MalformedError):
        return f"the document is damaged or has no readable content ({detail})"
    if isinstance(error, anydoc.MissingPartError):
        return f"the document is incomplete ({detail})"
    if isinstance(error, anydoc.ResourceLimitError):
        return f"the document crosses a safety limit and was not converted ({detail})"
    return f"the document could not be converted ({detail})"


def _anydoc_to_markdown(path: Path) -> str:
    # Imported here, never at module level: a PDF must not pay for it, and the
    # base install does not have it.
    try:
        import anydoc
    except ImportError:
        raise ConversionError(_NEEDS_OFFICE_EXTRA) from None
    try:
        # ocr="reject" is the privacy guarantee, not a default to be tidied up.
        # anydoc's other mode, "hosted", uploads the whole document to
        # Firecrawl; with an api_key or api_url it would do so quietly. Nothing
        # here passes either, and tests/test_convert_pure.py fails if anyone does.
        return anydoc.to_markdown(str(path), ocr="reject")
    except anydoc.ConvertError as e:
        raise ConversionError(_describe_anydoc_error(e)) from e


def convert_path(path: Path | str) -> ConversionResult:
    """Сконвертировать файл на диске. Синхронно: асинхронность -- забота веб-слоя."""
    path = Path(path)
    suffix = path.suffix.lower()

    pages = None
    try:
        with open(path, "rb") as f:
            head = f.read(1024)

        # PDFs take the checked pdf-inspector route (trimitdown.pdf_route). No
        # fallback to another PDF reader when the route refuses a file --
        # pdf-inspector, pypdf and pdfium were probed on damaged and encrypted
        # files and refuse together; a second reader would only turn a clear
        # refusal into an empty document. This is a route, not a fallback: a
        # file misnamed ".pdf" that is not a PDF at all (an HTML page saved with
        # the wrong extension) must still convert, so the %PDF magic bytes keep
        # real PDFs on our path and everything else goes to _convert_other,
        # which decides by content as well as by suffix.
        #
        # This check has now been written wrong twice. `data[:5] == b"%PDF"` was
        # always False (5-byte slice, 4-byte literal) and made the branch a
        # silent no-op. `data[:4] == b"%PDF"` looks right but anchors the marker
        # at byte 0 — real PDFs can have a leading \r\n or junk before the header
        # (measured: 2 of 719 real files, one offset by 2 bytes, one by 135) and
        # pdfplumber parses both fine. Anchoring routed them away from the PDF
        # route, which is exactly the defect this comment describes. pdfminer
        # itself scans for the marker rather than anchoring, which is why those
        # files open at all — so scan a window instead of anchoring at byte 0.
        if suffix == ".pdf" and head.find(b"%PDF") != -1:
            pages = tuple(route_pdf(path))
            text = render_markdown(list(pages))
        else:
            text = _convert_other(path, suffix, head)
    except Exception as e:
        raise ConversionError(str(e)) from e

    before, unit, units = _estimate_before_tokens(suffix, str(path))

    # Token count is a non-essential stat — the converted document must always be
    # returned. tiktoken can fail in a packaged build (e.g. the tiktoken_ext
    # encoding-constructor plugin not bundled), and an unguarded call there turned
    # every local-mode conversion into a 500. Degrade to a null count instead.
    try:
        after = count_tokens(text)
    except Exception:
        after = None

    return ConversionResult(
        text=text, tokens_after=after, tokens_before=before, unit=unit, units=units,
        pages=pages,
    )


def convert_bytes(data: bytes, suffix: str) -> ConversionResult:
    """Сконвертировать содержимое в памяти.

    Существует ради веб-слоя, где вход -- UploadFile, а не путь. Временный файл
    нужен потому, что и anydoc, и маршрут PDF работают с путями.
    """
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        return convert_path(tmp_path)
    finally:
        Path(tmp_path).unlink(missing_ok=True)
